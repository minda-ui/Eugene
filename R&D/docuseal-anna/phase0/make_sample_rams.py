#!/usr/bin/env python3
"""
RND-8 Phase 0 — generate a tagged sample RAMS (.docx) for the DocuSeal trial test.

Purpose: a short, realistic RAMS carrying DocuSeal **text-tags** at the signature
blocks, so we can upload it to a free DocuSeal Cloud trial and confirm the fields
auto-detect — and specifically learn whether the DOCX-field-tags upload works on the
free/trial tier (the A-vs-B / free-vs-Pro question).

This mirrors how Anna builds RAMS (python-docx); the tags are just literal text.

DocuSeal tag syntax (confirmed 2026-10-11 from DocuSeal docs):
  {{Field Name;type=signature;role=Operative 1}}   -> signature for that signer
  {{Field Name;type=date;role=Operative 1}}        -> date the signer picks
  {{Field Name;type=datenow;role=Operative 1}}      -> auto-filled signing date (read-only)
  {{Field Name;role=Operative 1}}                   -> text field (no type = text)
  attributes are separated by semicolons; `role=` names the signer/submitter.

Run:  python3 make_sample_rams.py
Out:  Sample-RAMS-DocuSeal-Phase0.docx  (in this folder)

Guide-only test artifact. No secret held. Not a real RAMS — content is placeholder.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "Sample-RAMS-DocuSeal-Phase0.docx"

# The three signer roles DocuSeal will create from the tags below.
# When you create the submission you map each role -> a name + email.
ROLE_MD = "Director"          # issues/approves — e.g. Mindaugas
ROLE_OP1 = "Operative 1"
ROLE_OP2 = "Operative 2"


def tag(p, text, *, hide=False):
    """Add a DocuSeal text-tag run. hide=True makes it tiny/white so it doesn't
    show in the finished PDF (DocuSeal replaces the tag with the field either way;
    for the trial, leave visible so it's obvious where fields landed)."""
    run = p.add_run(text)
    if hide:
        run.font.size = Pt(1)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return run


doc = Document()

# ---- Title ------------------------------------------------------------------
h = doc.add_heading("RISK ASSESSMENT & METHOD STATEMENT (SAMPLE)", level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("DocuSeal Phase 0 test document — NOT a live RAMS")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic = True

meta = doc.add_paragraph()
meta.add_run("Project: ").bold = True
meta.add_run("FC-SAMPLE / DocuSeal trial     ")
meta.add_run("Document no: ").bold = True
meta.add_run("RAMS-SAMPLE-01     ")
meta.add_run("Rev: ").bold = True
meta.add_run("A")

# ---- A little body so it reads like a real doc ------------------------------
doc.add_heading("1. Task", level=1)
doc.add_paragraph(
    "Installation of shopfitting joinery within an occupied retail mall unit, "
    "including transport of fixtures across the mall and fixing to walls."
)
doc.add_heading("2. Key hazards & controls (extract)", level=1)
for hz, ctrl in [
    ("Manual handling of fixtures", "Team lift, trolleys, planned route, no lone lifting."),
    ("Working at height (steps)", "Inspected steps, 3 points of contact, exclusion below."),
    ("Public interface in mall", "Barriers, out-of-hours deliveries, banksman across mall."),
    ("Power tools / trailing leads", "110V, PAT-tested, leads routed/covered, RCD protection."),
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(hz + ": ").bold = True
    p.add_run(ctrl)

doc.add_heading("3. Declaration & sign-off", level=1)

# ---- Issued/approved by (director) -----------------------------------------
doc.add_paragraph(
    "This RAMS has been reviewed and approved for issue. The signature below "
    "confirms director approval of the method and controls."
)
p = doc.add_paragraph()
p.add_run("Issued & approved by (Director): ").bold = True
p.add_run("____________________    ")
p.add_run("Signature: ")
tag(p, f"{{{{Director Signature;type=signature;role={ROLE_MD}}}}}")
p2 = doc.add_paragraph()
p2.add_run("Name: Mindaugas Gaudiesius      ")
p2.add_run("Date: ")
tag(p2, f"{{{{Director Date;type=datenow;role={ROLE_MD}}}}}")

doc.add_paragraph()  # spacer

# ---- Operative sign-off sheet ----------------------------------------------
doc.add_paragraph(
    "I have read and understood this RAMS and will work in accordance with it:"
).runs[0].bold = True

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
for i, t in enumerate(["Operative (print name)", "Signature", "Date", "Company"]):
    hdr[i].paragraphs[0].add_run(t).bold = True

for role in (ROLE_OP1, ROLE_OP2):
    cells = table.add_row().cells
    # print-name = a plain text field the signer fills
    tag(cells[0].paragraphs[0], f"{{{{{role} Name;role={role}}}}}")
    tag(cells[1].paragraphs[0], f"{{{{{role} Signature;type=signature;role={role}}}}}")
    tag(cells[2].paragraphs[0], f"{{{{{role} Date;type=date;role={role}}}}}")
    tag(cells[3].paragraphs[0], f"{{{{{role} Company;role={role}}}}}")

doc.add_paragraph()
note = doc.add_paragraph(
    "Phase 0 note: the {{...}} markers above are DocuSeal text-tags. On upload DocuSeal "
    "should auto-detect them as fields for the Director, Operative 1 and Operative 2 "
    "submitters. If the upload rejects the tags or ignores them, the DOCX-field-tags API "
    "is Pro-gated (answers free-vs-Pro)."
)
note.runs[0].italic = True
note.runs[0].font.size = Pt(9)

doc.save(OUT)
print(f"wrote {OUT}")
print("Submitters DocuSeal will expect (map each to a name+email at submission time):")
for r in (ROLE_MD, ROLE_OP1, ROLE_OP2):
    print(f"  - {r}")
